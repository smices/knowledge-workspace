from io import BytesIO

from docx import Document

from worker.main import chunks, extract


def test_docx_headings_become_chunk_section_paths():
    document = Document()
    document.add_heading("审批制度", level=1)
    document.add_paragraph("预算申请需要部门负责人审批。")
    buffer = BytesIO()
    document.save(buffer)
    items = chunks(extract(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    assert items == [("预算申请需要部门负责人审批。", ["审批制度"], None)]
